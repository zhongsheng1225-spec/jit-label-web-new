from __future__ import annotations

import csv
import io
import re
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
import streamlit as st
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


APP_DIR = Path(__file__).resolve().parent
MAPPING_XLSX = APP_DIR / "data" / "mapping.xlsx"
SKU_RE = re.compile(r"\b(?:[A-Z]\d{7,12}|\d{10,14})(?:-\d+)?\b")
DEFAULT_FONT = "STSong-Light"


def safe_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def pick_mapping_sheet(workbook: openpyxl.Workbook):
    if "Sheet2" in workbook.sheetnames:
        return workbook["Sheet2"]
    if not workbook.sheetnames:
        raise ValueError("Excel 中没有可读取的工作表。")
    return workbook[workbook.sheetnames[0]]


def mapping_start_row(ws) -> int:
    first_code = safe_text(ws.cell(row=1, column=1).value)
    first_name = safe_text(ws.cell(row=1, column=2).value)
    if "商品编码" in first_code or "商品名称" in first_name:
        return 2
    return 1


@st.cache_data(show_spinner=False)
def load_mapping_excel(path: str) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = pick_mapping_sheet(wb)
    mapping: dict[str, str] = {}
    for code_raw, name_raw in ws.iter_rows(min_row=mapping_start_row(ws), min_col=1, max_col=2, values_only=True):
        code = safe_text(code_raw)
        name = safe_text(name_raw)
        if code and name:
            mapping[code] = name
    if not mapping:
        raise ValueError("Excel 中没有读取到有效的商品编码和商品名称。")
    return mapping


def save_uploaded_pdf(uploaded_file) -> Path:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(uploaded_file.getvalue())
        return Path(tmp.name)


def annotated_download_name(uploaded_name: str) -> str:
    original = Path(uploaded_name or "JIT面单.pdf").stem
    return f"{original}注释.pdf"


def extract_code_from_page(page: pdfplumber.page.Page) -> str:
    text = page.extract_text() or ""
    matches = SKU_RE.findall(text)
    letter_codes = [match for match in matches if re.match(r"^[A-Z]", match)]
    if letter_codes:
        return letter_codes[0]
    numeric_codes = [
        match
        for match in matches
        if re.match(r"^\d{10,14}(?:-\d+)?$", match) and not match.startswith("120000")
    ]
    return numeric_codes[0] if numeric_codes else ""


def is_pickup_page(text: str) -> bool:
    pickup_markers = ["入库单号", "仓库编码", "SKU数量", "物流单号"]
    if sum(marker in text for marker in pickup_markers) >= 2:
        return True
    # Some carrier PDFs expose the Chinese labels as garbled text, but keep
    # stable order identifiers such as PONY/AP and omit the product-label line.
    return "Product Name" not in text and ("PONY" in text or "TRAN_STORE" in text)


def product_info(raw_code: str, mapping: dict[str, str]) -> dict[str, str]:
    lookup_code = raw_code.split("-", 1)[0] if raw_code else ""
    name = mapping.get(raw_code) or mapping.get(lookup_code) or ""
    status = "通过"
    notes: list[str] = []

    if not raw_code:
        status = "需处理"
        notes.append("货品标签页未识别到商品编码")
    elif raw_code != lookup_code and lookup_code in mapping:
        notes.append(f"标签编码为 {raw_code}，已按主编码 {lookup_code} 匹配")
    if raw_code and not name:
        status = "需处理"
        notes.append("对应表未找到商品名称")

    return {
        "raw_code": raw_code,
        "lookup_code": lookup_code,
        "name": name,
        "status": status,
        "notes": "；".join(notes),
    }


def aggregate_products(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    aggregated: list[dict[str, Any]] = []
    by_code: dict[str, dict[str, Any]] = {}
    for product in products:
        key = product.get("lookup_code") or product.get("raw_code") or f"unrecognized-{len(aggregated)}"
        existing = by_code.get(key)
        if existing is None:
            existing = {
                "lookup_code": product.get("lookup_code", ""),
                "raw_code": product.get("raw_code", ""),
                "name": product.get("name", ""),
                "quantity": 0,
                "status": product.get("status", "需处理"),
            }
            by_code[key] = existing
            aggregated.append(existing)
        existing["quantity"] += 1
        if product.get("status") != "通过":
            existing["status"] = "需处理"
    return aggregated


def analyze_pdf(pdf_path: Path, mapping: dict[str, str]) -> dict[str, Any]:
    reader = PdfReader(str(pdf_path))
    page_count = len(reader.pages)
    errors: list[str] = []
    warnings: list[str] = []

    if page_count == 0:
        errors.append("PDF 没有页面。")

    sizes = [
        (round(float(page.mediabox.width), 2), round(float(page.mediabox.height), 2))
        for page in reader.pages
    ]
    unique_sizes = sorted(set(sizes))
    if len(unique_sizes) > 1:
        warnings.append("PDF 内页面尺寸不完全一致；系统会按每组揽收面单尺寸插入注释页。")

    groups: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    seen_codes: dict[str, int] = {}

    with pdfplumber.open(str(pdf_path)) as pdf:
        page_texts = [page.extract_text() or "" for page in pdf.pages]
        pickup_indexes = [idx for idx, text in enumerate(page_texts) if is_pickup_page(text)]

        if not pickup_indexes:
            errors.append("未识别到揽收面单页，请检查 PDF。")

        output_page_cursor = 1
        for group_idx, pickup_index in enumerate(pickup_indexes):
            next_pickup = pickup_indexes[group_idx + 1] if group_idx + 1 < len(pickup_indexes) else page_count
            label_indexes = [
                idx
                for idx in range(pickup_index + 1, next_pickup)
                if "Product Name" in page_texts[idx] or extract_code_from_page(pdf.pages[idx])
            ]

            products: list[dict[str, Any]] = []
            for label_index in label_indexes:
                raw_code = extract_code_from_page(pdf.pages[label_index])
                info = product_info(raw_code, mapping)
                if info["lookup_code"]:
                    seen_codes[info["lookup_code"]] = seen_codes.get(info["lookup_code"], 0) + 1
                product = {
                    "label_page_index": label_index,
                    "label_page": label_index + 1,
                    **info,
                }
                products.append(product)
                rows.append(
                    {
                        "组号": group_idx + 1,
                        "揽收页": pickup_index + 1,
                        "标签页": label_index + 1,
                        "标签编码": info["raw_code"],
                        "匹配编码": info["lookup_code"],
                        "商品名称": info["name"],
                        "状态": info["status"],
                        "备注": info["notes"],
                        "输出注释页": output_page_cursor,
                    }
                )

            if not products:
                rows.append(
                    {
                        "组号": group_idx + 1,
                        "揽收页": pickup_index + 1,
                        "标签页": "",
                        "标签编码": "",
                        "匹配编码": "",
                        "商品名称": "",
                        "状态": "需处理",
                        "备注": "该揽收面单后未识别到产品标签",
                        "输出注释页": output_page_cursor,
                    }
                )

            group_end = next_pickup - 1
            groups.append(
                {
                    "group": group_idx + 1,
                    "pickup_page_index": pickup_index,
                    "pickup_page": pickup_index + 1,
                    "end_page_index": group_end,
                    "products": products,
                    "display_products": aggregate_products(products),
                    "output_note_page": output_page_cursor,
                }
            )
            output_page_cursor += (group_end - pickup_index + 1) + 1

    repeated = sorted(code for code, count in seen_codes.items() if count > 1)
    if repeated:
        warnings.append("同一个商品编码在本 PDF 中出现多次：" + "、".join(repeated))
    if any(row["状态"] != "通过" for row in rows):
        errors.append("有订单未通过复核，请先更新对应表或检查 PDF。")

    return {
        "page_count": page_count,
        "order_count": len(groups),
        "expected_output_pages": page_count + len(groups),
        "errors": errors,
        "warnings": warnings,
        "groups": groups,
        "rows": rows,
        "passed": not errors,
    }


def register_font() -> str:
    try:
        pdfmetrics.getFont(DEFAULT_FONT)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(DEFAULT_FONT))
    return DEFAULT_FONT


def wrap_text(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        trial = current + char
        if pdfmetrics.stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
            continue
        if current:
            lines.append(current)
        current = char
    if current:
        lines.append(current)
    return lines


def make_note_page(width: float, height: float, products: list[dict[str, Any]]) -> PdfReader:
    font_name = register_font()
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=(width, height))
    c.setFillColor(HexColor("#FFFFFF"))
    c.rect(0, 0, width, height, fill=1, stroke=0)

    margin = 12
    max_width = width - margin * 2
    title_size = 24
    body_size = 20
    label_size = 14
    product_lines: list[str] = []
    for idx, product in enumerate(products, 1):
        code = product.get("lookup_code") or product.get("raw_code") or "未识别"
        name = product.get("name") or "未找到商品名称"
        quantity = product.get("quantity", 1)
        quantity_text = f" x {quantity}" if quantity > 1 else ""
        product_lines.extend(wrap_text(f"{idx}. {code}{quantity_text}", font_name, body_size, max_width))
        product_lines.extend(wrap_text(name, font_name, body_size, max_width))
    while len(product_lines) > 9 and body_size > 11:
        body_size -= 0.4
        product_lines = []
        for idx, product in enumerate(products, 1):
            code = product.get("lookup_code") or product.get("raw_code") or "未识别"
            name = product.get("name") or "未找到商品名称"
            quantity = product.get("quantity", 1)
            quantity_text = f" x {quantity}" if quantity > 1 else ""
            product_lines.extend(wrap_text(f"{idx}. {code}{quantity_text}", font_name, body_size, max_width))
            product_lines.extend(wrap_text(name, font_name, body_size, max_width))

    y = height - 20
    c.setFillColor(HexColor("#111111"))
    c.setFont(font_name, title_size)
    c.drawString(margin, y, "商品名称注释")

    y -= 21
    c.setFont(font_name, label_size)
    c.setFillColor(HexColor("#555555"))
    total_quantity = sum(product.get("quantity", 1) for product in products)
    c.drawString(margin, y, f"商品种类：{len(products)}  总件数：{total_quantity}")
    y -= 15
    c.setFont(font_name, body_size)
    c.setFillColor(HexColor("#111111"))
    for line in product_lines:
        if y < margin:
            break
        c.drawString(margin, y, line)
        y -= body_size + 2.4

    c.showPage()
    c.save()
    buffer.seek(0)
    return PdfReader(buffer)


def create_annotated_pdf(pdf_path: Path, review: dict[str, Any]) -> bytes:
    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()
    group_by_pickup = {group["pickup_page_index"]: group for group in review["groups"]}
    for page_index, page in enumerate(reader.pages):
        group = group_by_pickup.get(page_index)
        if group:
            note_pdf = make_note_page(
                float(page.mediabox.width),
                float(page.mediabox.height),
                group["display_products"],
            )
            writer.add_page(note_pdf.pages[0])
        writer.add_page(page)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def review_csv_bytes(review: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    fieldnames = ["组号", "揽收页", "标签页", "标签编码", "匹配编码", "商品名称", "状态", "备注", "输出注释页"]
    writer = csv.DictWriter(buffer, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(review["rows"])
    return buffer.getvalue().encode("utf-8-sig")


st.set_page_config(page_title="JIT 面单商品名称注释", page_icon="📦", layout="wide")
st.title("JIT 面单商品名称注释")
st.caption("上传速卖通 JIT 面单 PDF，系统会先复核商品编码，再生成已插入商品名称注释页的新 PDF。")

try:
    mapping = load_mapping_excel(str(MAPPING_XLSX))
except Exception as exc:
    mapping = {}
    st.error(f"对应表读取失败：{exc}")

with st.sidebar:
    st.header("对应表")
    if mapping:
        st.success(f"当前对应表：mapping.xlsx，{len(mapping)} 个编码")
    else:
        st.error("当前没有可用对应表")
    st.write("长期更新：替换 GitHub 里的 data/mapping.xlsx。A 列是商品编码，B 列是商品名称。")

uploaded_pdf = st.file_uploader("面单 PDF", type=["pdf"])

if not uploaded_pdf:
    st.info("请先上传面单 PDF。")
    st.stop()

if not mapping:
    st.stop()

pdf_path = save_uploaded_pdf(uploaded_pdf)
try:
    with st.spinner("正在复核商品编码和对应表"):
        review = analyze_pdf(pdf_path, mapping)
except Exception as exc:
    st.error(f"PDF 读取失败：{exc}")
    st.stop()

left, right = st.columns(2)
left.metric("源 PDF 页数", review["page_count"])
right.metric("识别订单数", review["order_count"])

for warning in review["warnings"]:
    st.warning(warning)
for error in review["errors"]:
    st.error(error)
if review["passed"]:
    st.success(f"复核通过，预计输出 {review['expected_output_pages']} 页。")

st.subheader("复核清单")
st.dataframe(review["rows"], use_container_width=True, hide_index=True)

st.download_button(
    "下载复核清单 CSV",
    review_csv_bytes(review),
    file_name="JIT面单复核清单.csv",
    mime="text/csv",
)

if review["passed"]:
    with st.spinner("正在生成 PDF"):
        output_pdf = create_annotated_pdf(pdf_path, review)
    st.download_button(
        "下载已加商品名称注释的 PDF",
        output_pdf,
        file_name=annotated_download_name(uploaded_pdf.name),
        mime="application/pdf",
        type="primary",
    )
