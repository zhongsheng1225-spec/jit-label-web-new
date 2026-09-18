from __future__ import annotations

import csv
import io
import re
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


SKU_RE = re.compile(r"\b(?:[A-Z]\d{7,12}|\d{10,14})(?:-\d+)?\b")
DEFAULT_FONT = "STSong-Light"


def _register_font() -> str:
    try:
        pdfmetrics.getFont(DEFAULT_FONT)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(DEFAULT_FONT))
    return DEFAULT_FONT


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_mapping_csv(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        for row in reader:
            code = safe_text(row.get("商品编码") or row.get("code"))
            name = safe_text(row.get("商品名称") or row.get("name"))
            if code and name:
                mapping[code] = name
    return mapping


def excel_to_mapping_csv_bytes(excel_bytes: bytes) -> tuple[bytes, int]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(excel_bytes)
        tmp_path = Path(tmp.name)
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True, read_only=True)
        ws = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb[wb.sheetnames[1]]
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["商品编码", "商品名称"])
        count = 0
        for code_raw, name_raw in ws.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True):
            code = safe_text(code_raw)
            name = safe_text(name_raw)
            if not code or not name:
                continue
            writer.writerow([code, name])
            count += 1
        return buffer.getvalue().encode("utf-8-sig"), count
    finally:
        tmp_path.unlink(missing_ok=True)


def _extract_code_from_page(page: pdfplumber.page.Page) -> str:
    text = page.extract_text() or ""
    matches = SKU_RE.findall(text)
    letter_codes = [match for match in matches if re.match(r"^[A-Z]", match)]
    if letter_codes:
        return letter_codes[0]

    numeric_codes = [
        match
        for match in matches
        if re.match(r"^\d{10,14}(?:-\d+)?$", match) and not match.startswith("120000")
    ]
    return numeric_codes[0] if numeric_codes else ""


def analyze_pdf(pdf_path: Path, mapping: dict[str, str]) -> dict[str, Any]:
    reader = PdfReader(str(pdf_path))
    page_count = len(reader.pages)
    errors: list[str] = []
    warnings: list[str] = []

    if page_count == 0:
        errors.append("PDF 没有页面。")
    if page_count % 2 != 0:
        errors.append(f"PDF 页数为 {page_count}，不是偶数；需要按“揽收面单 + 货品标签”两页一组处理。")

    sizes = [
        (round(float(page.mediabox.width), 2), round(float(page.mediabox.height), 2))
        for page in reader.pages
    ]
    unique_sizes = sorted(set(sizes))
    if len(unique_sizes) > 1:
        warnings.append("PDF 内页面尺寸不完全一致；系统会按每组揽收面单尺寸插入注释页。")

    rows: list[dict[str, Any]] = []
    order_count = page_count // 2
    seen_codes: dict[str, int] = {}

    with pdfplumber.open(str(pdf_path)) as pdf:
        for order_idx in range(order_count):
            label_page_index = order_idx * 2 + 1
            raw_code = _extract_code_from_page(pdf.pages[label_page_index])
            lookup_code = raw_code.split("-", 1)[0] if raw_code else ""
            name = mapping.get(raw_code) or mapping.get(lookup_code) or ""
            status = "通过"
            notes: list[str] = []

            if not raw_code:
                status = "需处理"
                notes.append("货品标签页未识别到商品编码")
            elif raw_code != lookup_code and lookup_code in mapping:
                notes.append(f"标签编码为 {raw_code}，已按主编码 {lookup_code} 匹配")
            if raw_code and not name:
                status = "需处理"
                notes.append("对应表未找到商品名称")
            if lookup_code:
                seen_codes[lookup_code] = seen_codes.get(lookup_code, 0) + 1

            rows.append(
                {
                    "订单": order_idx + 1,
                    "揽收页": order_idx * 2 + 1,
                    "标签页": label_page_index + 1,
                    "标签编码": raw_code,
                    "匹配编码": lookup_code,
                    "商品名称": name,
                    "状态": status,
                    "备注": "；".join(notes),
                    "输出注释页": order_idx * 3 + 1,
                }
            )

    repeated = sorted(code for code, count in seen_codes.items() if count > 1)
    if repeated:
        warnings.append("同一个商品编码在本 PDF 中出现多次：" + "、".join(repeated))

    if any(row["状态"] != "通过" for row in rows):
        errors.append("有订单未通过复核，请先更新对应表或检查 PDF。")

    return {
        "page_count": page_count,
        "order_count": order_count,
        "unique_page_sizes": unique_sizes,
        "expected_output_pages": order_count * 3,
        "errors": errors,
        "warnings": warnings,
        "rows": rows,
        "passed": not errors,
    }


def _wrap_text(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        trial = current + char
        if pdfmetrics.stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
            continue
        if current:
            lines.append(current)
        current = char
    if current:
        lines.append(current)
    return lines


def _make_note_page(width: float, height: float, code: str, name: str) -> PdfReader:
    font_name = _register_font()
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=(width, height))
    c.setFillColor(HexColor("#FFFFFF"))
    c.rect(0, 0, width, height, fill=1, stroke=0)

    margin = 12
    max_width = width - margin * 2
    title_size = 13
    body_size = 10.5
    label_size = 8.5

    name_lines = _wrap_text(name, font_name, body_size, max_width)
    while len(name_lines) > 8 and body_size > 5.8:
        body_size -= 0.4
        name_lines = _wrap_text(name, font_name, body_size, max_width)

    y = height - 18
    c.setFillColor(HexColor("#111111"))
    c.setFont(font_name, title_size)
    c.drawString(margin, y, "商品名称注释")

    y -= 18
    c.setFont(font_name, label_size)
    c.setFillColor(HexColor("#555555"))
    c.drawString(margin, y, "商品编码")
    y -= 11
    c.setFont(font_name, 10.5)
    c.setFillColor(HexColor("#111111"))
    c.drawString(margin, y, code or "未识别")

    y -= 18
    c.setFont(font_name, label_size)
    c.setFillColor(HexColor("#555555"))
    c.drawString(margin, y, "商品名称")
    y -= 12
    c.setFont(font_name, body_size)
    c.setFillColor(HexColor("#111111"))
    line_height = body_size + 2.2
    for line in name_lines:
        if y < margin:
            break
        c.drawString(margin, y, line)
        y -= line_height

    c.showPage()
    c.save()
    buffer.seek(0)
    return PdfReader(buffer)


def create_annotated_pdf(pdf_path: Path, review: dict[str, Any]) -> bytes:
    if not review["passed"]:
        raise ValueError("复核未通过，不能生成 PDF。")

    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()

    for row in review["rows"]:
        pickup_index = (row["订单"] - 1) * 2
        pickup_page = reader.pages[pickup_index]
        width = float(pickup_page.mediabox.width)
        height = float(pickup_page.mediabox.height)
        note_pdf = _make_note_page(width, height, row["匹配编码"], row["商品名称"])
        writer.add_page(note_pdf.pages[0])
        writer.add_page(pickup_page)
        writer.add_page(reader.pages[pickup_index + 1])

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def review_csv_bytes(review: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    fieldnames = ["订单", "揽收页", "标签页", "标签编码", "匹配编码", "商品名称", "状态", "备注", "输出注释页"]
    writer = csv.DictWriter(buffer, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(review["rows"])
    return buffer.getvalue().encode("utf-8-sig")
